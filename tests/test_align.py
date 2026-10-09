"""Offline tests for the kwartaalplanbord ↔ monday.com alignment and its app."""

from __future__ import annotations

import json
import re
from collections.abc import Iterator
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from starlette.testclient import TestClient

from builders import board_item
from mondaycom import align, align_export, planning
from mondaycom.config import DAM, DISTRIBUTION_BOARD, EPIC_BOARD, EPIC_GROUP_ACTIVE, EPIC_GROUP_BACKLOG, NON_DAM
from mondaycom.lookups import Choice

HTMX = {"HX-Request": "1"}

DUMP: dict[str, Any] = {
    "config": {
        "main": {
            "epics": [
                {"id": "e2", "name": "Maandrapportage MT", "dpr": "DPR-227", "portfolio": "FIR"},
                {"id": "e5", "name": "Personeellasten", "dpr": "DPR-227"},
                {"id": "e3", "name": "Spendanalyse", "dpr": "dpr-307"},
                {"id": "e1", "name": "Dashboard HWBP"},
            ],
            "sprints": [{"id": "s4", "end": "2026-12-25"}, {"id": "s1", "end": "2026-10-23"}],
        }
    },
    "people": {"agnes": {"name": "Agnes"}, "jelle": {"name": "Jelle"}},
    "stories": {
        "a": {"epicId": "e2", "pts": {"vis": 2}},
        "b": {"epicId": "e2", "pts": {"mod": 1}},
        "c": {"epicId": "e5", "pts": {"vis": 3}},
        "d": {"epicId": "e3", "pts": {"mod": 11, "vis": 4, "ia": 0}},
        "e": {"epicId": "e1", "pts": {"ds": 2}},
    },
}


def split(id: str, epic_id: str, todo: float, *shares: float | None) -> planning.Split:
    return planning.Split(
        id=id,
        name=id,
        epic_id=epic_id,
        todo=todo,
        has_todo=True,
        shares=dict(zip(planning.DISCIPLINES, shares, strict=True)),
    )


def monday_epic(id: str, prj_nr: str, todo: float, group: str = EPIC_GROUP_ACTIVE, **kw: Any) -> align.MondayEpic:
    plan = planning.PlanEpic(id=id, name=f"epic {prj_nr}", prj_nr=prj_nr, group=group, **kw)
    return align.MondayEpic(plan, split=split(f"s{id}", id, todo, 35, 30, 25, 10))


def monday() -> align.Monday:
    return align.Monday(
        epics=[
            monday_epic("227", "DPR-227", 9),
            monday_epic("307", "DPR-307", 0, group=EPIC_GROUP_BACKLOG),
            monday_epic("400", "DPR-400", 20),  # promised, not on the board
            monday_epic("401", "DPR-401", 20, due=date(2027, 3, 1)),  # later — not this quarter
            monday_epic("402", "DPR-402", 0),  # promised, nothing left
        ],
        people=[
            Choice(id="1", name="Agnes Dubbink", email="agnes@example.org"),
            Choice(id="2", name="Jelle de Jong", email="j@example.org"),
        ],
        current_end=date(2026, 10, 25),
    )


def test_split_from_whole_percentages_that_add_up() -> None:
    assert align.split_from({"DE": 1, "DB": 2}) == {"DE": 33, "DB": 67, "DS": 0, "PO/AT": 0}
    assert align.split_from({"DE": 1, "DB": 1, "DS": 1}) == {"DE": 34, "DB": 33, "DS": 33, "PO/AT": 0}
    assert align.split_from({}) == {}
    for points in ({"DE": 13, "DS": 26}, {"DE": 6, "DB": 1, "DS": 8, "PO/AT": 5}, {"DB": 3, "PO/AT": 16}):
        assert sum(align.split_from(points).values()) == 100


def test_parse_board_maps_disciplines_and_sums_stories() -> None:
    board = align.parse_board(DUMP)
    by_id = {e.id: e for e in board.epics}
    assert by_id["e2"].points == {"DB": 2.0, "DE": 1.0}
    assert by_id["e3"].dpr == "DPR-307"
    assert by_id["e3"].points == {"DE": 11.0, "DB": 4.0}  # a zero leaves no key
    assert board.quarter_end == date(2026, 12, 31)
    assert board.people == ["Agnes", "Jelle"]


def test_parse_board_refuses_an_unknown_discipline() -> None:
    bad = {**DUMP, "stories": {"x": {"epicId": "e2", "pts": {"pm": 1}}}}
    with pytest.raises(ValueError, match="pm"):
        align.parse_board(bad)


def test_compare_links_on_dpr_and_sums_a_shared_one() -> None:
    rows = align.compare(align.parse_board(DUMP), monday(), align.State())
    by_key = {r.key: r for r in rows}
    shared = by_key["DPR-227"]
    assert [e.id for e in shared.board] == ["e2", "e5"]
    assert shared.board_total == 6
    assert any("2 epics" in w for w in shared.warnings)
    assert by_key["db:e1"].kind == align.BOARD_ONLY
    # Only monday.com's promised epics with work left join as its own rows.
    assert [r.key for r in rows if r.kind == align.MONDAY_ONLY] == ["DPR-400"]


def test_a_relink_moves_an_epic_out_of_a_dpr() -> None:
    state = align.State(links={"e5": ""})
    by_key = {r.key: r for r in align.compare(align.parse_board(DUMP), monday(), state)}
    assert by_key["DPR-227"].board_split == {"DE": 33, "DB": 67, "DS": 0, "PO/AT": 0}
    assert by_key["db:e5"].kind == align.BOARD_ONLY


def test_the_selection_is_the_planning_pages() -> None:
    board = align.parse_board(DUMP)
    end = date(2026, 12, 31)

    def keys(sel: align.Selection) -> dict[str, bool]:
        return {r.key: r.selected for r in align.compare(board, monday(), align.State(), sel)}

    promised = keys(align.Selection(end))
    assert promised["DPR-227"] and not promised["DPR-307"]  # 307 is in the Backlog group
    assert "DPR-401" not in promised  # due after the quarter: Later
    later = keys(align.Selection(end, layers=(planning.LATER,)))
    assert "DPR-401" in later and "DPR-400" not in later
    # "Dit kwartaal" drops what has no due date; none of these has one.
    assert not keys(align.Selection(end, this_quarter=True))["DPR-227"]
    # A board epic linked to nothing keeps only the DAM filter, on its own portfolio name.
    assert not keys(align.Selection(end, dam=DAM))["db:e1"]
    assert keys(align.Selection(end, dam=NON_DAM))["db:e1"]


def test_flagged_past_ten_percent_of_monday() -> None:
    def row(board: float, todo: float) -> align.Row:
        e = align.BoardEpic(id="x", name="x", points={"DE": board})
        return align.Row(key="k", dpr="k", monday=monday_epic("1", "k", todo), board=(e,))

    assert not row(11, 10).flagged
    assert row(11.5, 10).flagged
    assert row(8, 10).flagged
    assert row(3, 0).flagged and row(3, 0).ratio == float("inf")
    assert not row(0, 0).flagged


def test_split_changes_and_backup_then_restore(tmp_path: Path) -> None:
    state = align.State(links={"e5": ""})
    rows = align.compare(align.parse_board(DUMP), monday(), state)
    changes = align.split_changes(rows)
    assert {r.dpr for r in changes} == {"DPR-227", "DPR-307"}

    sent: list[dict[str, Any]] = []

    class Client:
        def execute(self, query: str, variables: dict[str, Any] | None = None) -> dict[str, Any]:
            sent.append(variables or {})
            return {}

    backup = tmp_path / "backup.json"
    written = align.apply_splits(Client(), changes, backup)  # type: ignore[arg-type]
    assert json.loads(backup.read_text())["s227"]["shares"]["DE"] == 35
    values = json.loads(sent[0]["values"])
    assert values[DISTRIBUTION_BOARD.column("DE")] == "33"
    assert values[DISTRIBUTION_BOARD.column("DB")] == "67"
    assert written[0].shares["DB"] == 67

    sent.clear()
    assert align.restore_splits(Client(), backup) == 2  # type: ignore[arg-type]
    assert json.loads(sent[0]["values"])[DISTRIBUTION_BOARD.column("DE")] == "35"


def test_state_round_trip(tmp_path: Path) -> None:
    state = align.State(path=tmp_path / "s.json")
    state.decisions["DPR-227"] = align.Decision(wanted=6, note="zo")
    action = state.add_action("DPR-227", " Navragen ", "Agnes")
    state.set_status(action.id, "verstuurd")
    state.save()
    again = align.State.load(tmp_path / "s.json")
    assert again.decision("DPR-227").wanted == 6
    assert again.actions[0].text == "Navragen"
    assert again.actions[0].status == "verstuurd" and again.actions[0].sent


def test_message_carries_the_figures_and_the_note() -> None:
    state = align.State(links={"e5": ""})
    rows = {r.key: r for r in align.compare(align.parse_board(DUMP), monday(), state)}
    state.decisions["DPR-227"] = align.Decision(note="Dubbel geteld?")
    action = state.add_action("DPR-227", "Navragen", "Agnes Dubbink")
    text = align.message("Agnes Dubbink", [action], rows, state, "Jelle de Jong")
    assert text.startswith("Hoi Agnes,")
    assert "Kwartaalplanbord 3 STP · monday.com 9 STP" in text
    assert "Definition of Done: Dubbel geteld?" in text


def test_monday_epic_is_the_planning_epic_with_its_project_number() -> None:
    item = board_item(
        EPIC_BOARD,
        name="E",
        id="9",
        group={"id": EPIC_GROUP_ACTIVE, "title": "Actief"},
        prj_nr="DPR-9",
        due_date="2026-11-01",
    )
    e = align.MondayEpic(planning.PlanEpic.from_item(item))
    assert (e.prj_nr, e.group_title, e.plan.due) == ("DPR-9", "Actief", date(2026, 11, 1))
    assert e.plan.layer(date(2026, 12, 31)) == planning.PROMISED


# --- the app -----------------------------------------------------------------------------


@pytest.fixture
def app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[TestClient]:
    dump = tmp_path / "board.json"
    dump.write_text(json.dumps(DUMP))
    monkeypatch.setenv("ALIGN_BOARD", str(dump))
    monkeypatch.setenv("ALIGN_STATE", str(tmp_path / "state.json"))
    from mondaycom import align_web

    monkeypatch.setattr(align_web, "_MONDAY", [monday()])
    monkeypatch.setattr(align, "fetch_monday", lambda client: pytest.fail("no network in tests"))
    # As the app's own page asks: on its own host name, and through htmx.
    yield TestClient(align_web.app, base_url="http://127.0.0.1", headers={"HX-Request": "true"})


def test_comparison_marks_the_far_apart_rows(app: TestClient) -> None:
    html = app.get("/vergelijking?layer=promised&layer=backlog", headers=HTMX).text
    assert 'id="rij-DPR-227"' in html
    assert 'id="rij-db-e1"' in html
    assert html.count('class="flag"') == 2  # DPR-227 (6 vs 9) and DPR-307 (15 vs 0)
    only = app.get("/vergelijking?soort=monday", headers=HTMX).text
    assert 'id="rij-DPR-400"' in only and 'id="rij-DPR-227"' not in only


def test_the_planning_filters_scope_the_table(app: TestClient) -> None:
    html = app.get("/vergelijking", headers=HTMX).text  # Promised only, as on the Planning page
    assert 'id="rij-DPR-307"' not in html
    assert "1 epics uit het kwartaalplanbord (15 STP)" in html and "DPR-307 (Backlog)" in html
    assert 'value="2026-12-31"' in html and "hx-swap-oob" in html  # the quarter it settled on
    shown = app.get("/vergelijking?buiten=1", headers=HTMX).text
    assert 'id="rij-DPR-307"' in shown and "buiten de selectie" in shown
    later = app.get("/vergelijking?layer=later", headers=HTMX).text
    assert 'id="rij-DPR-401"' in later
    moved = app.get("/vergelijking?layer=later&end=2027-03-31", headers=HTMX).text
    assert 'id="rij-DPR-401"' not in moved  # due 2027-03-01: promised for that quarter
    dam = app.get(f"/vergelijking?dam={DAM}", headers=HTMX).text
    assert 'id="rij-db-e1"' not in dam


def test_the_page_carries_the_planning_filter_row(app: TestClient) -> None:
    html = app.get("/").text
    for word in ("Einde sprint", "Einde kwartaal", "Portfolio", "Toegezegd", "Dit kwartaal", "Herstellen"):
        assert word in html


def test_saving_a_decision_and_adding_an_action(app: TestClient, tmp_path: Path) -> None:
    reply = app.post("/bewaar", data={"key": "DPR-227", "wanted": "7,5", "note": "omdat"})
    assert "monday.com -1.5" in reply.text and "hx-swap-oob" in reply.text
    assert "aanpassen" in reply.text  # where to change it, deduced from the wanted total
    assert app.post("/actie", data={"key": "DPR-227", "tekst": "Vragen", "wie": "Agnes Dubbink"}).headers["HX-Refresh"]
    state = align.State.load(tmp_path / "state.json")
    assert state.decision("DPR-227").wanted == 7.5
    page = app.get("/acties").text
    assert "agnes@example.org" in page and "mailto:" in page and "Hoi Agnes," in page
    one = app.get("/acties?epic=DPR-227").text
    assert "Vragen" in one and "Nieuwe actie voor DPR-227" in one


def test_the_actions_column_counts_open_actions(tmp_path: Path) -> None:
    from fasthtml.common import to_xml

    from mondaycom import align_web

    state = align.State(path=tmp_path / "state.json")
    assert align_web.actions_count("DPR-227", state) == "-"
    action = state.add_action("DPR-227", "Navragen")
    state.add_action("DPR-227", "Nog iets")
    assert align_web.actions_count("DPR-227", state) == "2"
    state.set_status(action.id, "klaar")
    assert align_web.actions_count("DPR-227", state) == "1"
    for a in state.actions:
        state.set_status(a.id, "klaar")
    assert align_web.actions_count("DPR-227", state) == "0"
    assert "/acties?epic=DPR-227" in to_xml(align_web.actions_cell("DPR-227", state))


def test_where_follows_from_the_wanted_total() -> None:
    [row] = [r for r in align.compare(align.parse_board(DUMP), monday(), align.State()) if r.key == "DPR-227"]
    assert align.where(row, align.Decision()) == ""
    assert align.where(row, align.Decision(wanted=row.board_total)) in {"monday", "geen"}
    assert align.where(row, align.Decision(wanted=row.board_total + row.monday_total + 1)) == "beide"


def test_an_old_state_file_with_where_still_loads(tmp_path: Path) -> None:
    path = tmp_path / "state.json"
    path.write_text('{"decisions": {"k": {"wanted": 3, "where": "monday", "note": "x"}}}', encoding="utf-8")
    assert align.State.load(path).decision("k") == align.Decision(wanted=3, note="x")


def test_relinking_an_epic(app: TestClient, tmp_path: Path) -> None:
    reply = app.post("/koppel", data={"epic": "e1", "dpr": "400"})
    assert reply.headers["HX-Trigger"] == "refilter, gewijzigd"
    assert align.State.load(tmp_path / "state.json").links == {"e1": "DPR-400"}
    html = app.get("/vergelijking", headers=HTMX).text
    assert 'id="rij-db-e1"' not in html


def test_export_is_markdown(app: TestClient) -> None:
    app.post("/bewaar", data={"key": "DPR-227", "wanted": "6", "note": "ok"})
    md = app.get("/export").text
    assert "| DPR-227 |" in md and "| 6 |" in md


# --- saving a copy -----------------------------------------------------------------------


def saved_state(tmp_path: Path) -> tuple[align.State, list[align.Row]]:
    state = align.State(links={"e5": ""}, path=tmp_path / "state.json")
    state.decisions["DPR-227"] = align.Decision(wanted=6, note="Dubbel; zie actie")
    state.add_action("DPR-227", "Navragen", "Agnes Dubbink")
    rows = align.compare(align.parse_board(DUMP), monday(), state)
    return state, rows


def test_every_format_reads_back(tmp_path: Path) -> None:
    import csv
    import io

    from odf.opendocument import load
    from odf.table import Table
    from openpyxl import load_workbook

    state, rows = saved_state(tmp_path)
    book = load_workbook(io.BytesIO(align_export.render("xlsx", rows, state)))
    assert book.sheetnames == ["Afstemming", "Acties", "Koppelingen"]
    sheet = book["Afstemming"]
    head = [c.value for c in sheet[1]]
    line = next(r for r in sheet.iter_rows(min_row=2, values_only=True) if r[0] == "DPR-227")
    assert line[head.index("Gewenst STP")] == 6 and line[head.index("Definition of Done")] == "Dubbel; zie actie"
    assert line[head.index("Database STP")] == 3 and line[head.index("monday.com STP-TODO")] == 9
    assert book["Acties"]["D2"].value == "Navragen"

    (tmp_path / "x.ods").write_bytes(align_export.render("ods", rows, state))
    names = [t.getAttribute("name") for t in load(str(tmp_path / "x.ods")).spreadsheet.getElementsByType(Table)]
    assert names == ["Afstemming", "Acties", "Koppelingen"]

    text = align_export.render("csv", rows, state).decode("utf-8")
    assert text.startswith("\ufeffDPR;")
    parsed = list(csv.reader(io.StringIO(text.lstrip("\ufeff")), delimiter=";"))
    assert parsed[0][-1] == "Acties"
    assert any(r[0] == "DPR-227" and "Agnes Dubbink: Navragen (open)" in r[-1] for r in parsed)

    data = json.loads(align_export.render("json", rows, state))
    assert data["decisions"]["DPR-227"]["wanted"] == 6 and data["afstemming"][0]["DPR"]

    with pytest.raises(ValueError):
        align_export.render("pdf", rows, state)


def test_a_folder_must_exist_and_lie_outside_the_repository(tmp_path: Path) -> None:
    assert align_export.resolve_folder(f"  {tmp_path} ") == tmp_path
    for bad in ("", "relatief/pad", str(tmp_path / "nergens"), str(align.REPO), str(align.REPO / "docs" / "tmp")):
        with pytest.raises(ValueError):
            align_export.resolve_folder(bad)


def test_a_windows_path_is_read_as_its_wsl_mount() -> None:
    if not Path("/mnt/c").is_dir():
        pytest.skip("no WSL mount here")
    with pytest.raises(ValueError, match="/mnt/c/Nergens/Hier"):
        align_export.resolve_folder(r"C:\Nergens\Hier")


def test_unsaved_until_a_copy_holds_the_changes(tmp_path: Path) -> None:
    state = align.State(path=tmp_path / "s.json")
    assert not state.unsaved  # nothing decided, nothing to lose
    state.decisions["DPR-1"] = align.Decision(note="x")
    assert state.unsaved
    state.saving.folder = str(tmp_path)
    state.saving.format = "json"
    target = align_export.save_copy(state, [])
    assert target == tmp_path / "afstemming-kwartaalplanning.json" and target.exists()
    assert not state.unsaved
    state.saving.autosave = True  # where the copy goes is not a change to the decisions
    assert not state.unsaved
    state.decisions["DPR-1"] = align.Decision(note="y")
    assert state.unsaved
    state.save()
    assert align.State.load(tmp_path / "s.json").unsaved


def test_autosave_waits_for_its_interval(tmp_path: Path) -> None:
    from datetime import datetime, timedelta

    state = align.State(decisions={"k": align.Decision(note="x")})
    state.saving = align.Saving(folder=str(tmp_path), autosave=True, minutes=5)
    assert align_export.autosave_due(state)  # never saved
    align_export.record_saved(state, "ergens")
    state.decisions["k"] = align.Decision(note="y")
    now = datetime.fromisoformat(state.saving.saved_at)
    assert not align_export.autosave_due(state, now + timedelta(minutes=4))
    assert align_export.autosave_due(state, now + timedelta(minutes=5))
    state.saving.autosave = False
    assert not align_export.autosave_due(state, now + timedelta(hours=1))


def test_the_banner_warns_until_saved(app: TestClient, tmp_path: Path) -> None:
    assert "lokaal bewaard" in app.get("/opslag_status").text
    reply = app.post("/bewaar", data={"key": "DPR-227", "wanted": "6"})
    assert reply.headers["HX-Trigger"] == "gewijzigd"
    assert "Niet opgeslagen" in app.get("/opslag_status").text
    out = tmp_path / "eigen"
    out.mkdir()
    page = app.post("/opslaan_instellen", data={"format": "xlsx", "folder": str(out)})
    assert str(out / "afstemming-kwartaalplanning.xlsx") in page.text
    assert (out / "afstemming-kwartaalplanning.xlsx").exists()
    assert "Alles is opgeslagen" in app.get("/opslag_status").text
    app.post("/autosave_zetten", data={"autosave": "on", "minutes": "5"})
    app.post("/bewaar", data={"key": "DPR-227", "wanted": "7"})
    banner = app.get("/opslag_status").text
    assert "Niet opgeslagen" in banner and "automatisch" in banner and "Nu opslaan" in banner
    app.post("/nu_opslaan")
    assert "Alles is opgeslagen" in app.get("/opslag_status").text


def test_a_bad_folder_is_said_not_raised(app: TestClient, tmp_path: Path) -> None:
    reply = app.post("/opslaan_instellen", data={"format": "csv", "folder": str(align.REPO)})
    assert "openbaar" in reply.text
    assert "Opslaan lukte niet" in app.get("/opslag_status").text


def test_there_is_no_download(app: TestClient) -> None:
    assert app.get("/download?fmt=ods").status_code == 404
    assert "download" not in app.get("/opslaan").text.lower()


# --- only the app itself -------------------------------------------------------------------


def test_another_host_name_is_refused(app: TestClient) -> None:
    """DNS rebinding: another site's name pointing at 127.0.0.1 gets nothing back."""
    assert app.get("/vergelijking", headers={"Host": "evil.example"}).status_code == 400
    assert app.get("/vergelijking", headers={"Host": "localhost:5002"}).status_code == 200


def test_a_change_needs_post_and_the_apps_own_page(app: TestClient, tmp_path: Path) -> None:
    # An <img src=…> is a GET: the routes that change something do not answer it.
    assert app.get("/verdeling?key=DPR-227").status_code == 405
    assert app.get("/opslaan_instellen?folder=/tmp").status_code == 405
    # A cross-site form POST carries no HX-Request, a foreign Origin, or is marked cross-site.
    data = {"key": "DPR-227", "note": "geplant"}
    assert app.post("/bewaar", data=data, headers={"HX-Request": ""}).status_code == 403
    assert app.post("/bewaar", data=data, headers={"Origin": "https://evil.example"}).status_code == 403
    assert app.post("/bewaar", data=data, headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert app.get("/mappen", headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
    assert align.State.load(tmp_path / "state.json").decision("DPR-227").note == ""
    # The app's own page passes: same origin, through htmx.
    own = {"Origin": "http://127.0.0.1", "Sec-Fetch-Site": "same-origin"}
    assert app.post("/bewaar", data=data, headers=own).status_code == 200
    assert align.State.load(tmp_path / "state.json").decision("DPR-227").note == "geplant"


def test_verdeling_writes_only_a_split_that_differs(app: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from mondaycom import align_web

    written: list[str] = []
    monkeypatch.setattr(align, "apply_splits", lambda client, rows, backup: written.extend(r.key for r in rows) or [])
    monkeypatch.setattr(align_web, "monday_client", lambda: None)
    # DPR-400 is monday.com's alone: no board split, so nothing to write.
    assert app.post("/verdeling", data={"key": "DPR-400"}).status_code == 200
    assert written == []


# --- the folder dialog ---------------------------------------------------------------------


def test_browse_lists_subfolders_and_skips_hidden_ones(tmp_path: Path) -> None:
    for name in ("Planning", "archief", ".git", "$Recycle.Bin"):
        (tmp_path / name).mkdir()
    (tmp_path / "los.txt").write_text("x")
    listing = align_export.browse(str(tmp_path))
    assert [f.name for f in listing.folders] == ["archief", "Planning"]
    assert listing.problem == "" and listing.parent == tmp_path.parent


def test_a_missing_folder_opens_its_nearest_parent_with_the_name_ready(tmp_path: Path) -> None:
    listing = align_export.browse(str(tmp_path / "backup" / "align"))
    assert listing.path == tmp_path and listing.missing == "backup"


def test_make_folder_makes_it_and_goes_in(tmp_path: Path) -> None:
    listing = align_export.make_folder(str(tmp_path), " backup ")
    assert listing.path == tmp_path / "backup" and (tmp_path / "backup").is_dir() and not listing.error
    again = align_export.make_folder(str(tmp_path), "backup")  # already there: just go in
    assert again.path == tmp_path / "backup" and not again.error
    for bad in ("", "..", "a/b", "c:d", 'x"y'):
        refused = align_export.make_folder(str(tmp_path), bad)
        assert refused.path == tmp_path and "mapnaam" in refused.error


def test_the_repository_cannot_be_chosen_nor_built_in() -> None:
    listing = align_export.browse(str(align.REPO / "docs"))
    assert "openbaar" in listing.problem
    made = align_export.make_folder(str(align.REPO / "docs"), "kopie")
    assert "openbaar" in made.error and not (align.REPO / "docs" / "kopie").exists()


def test_windows_paths_read_both_ways() -> None:
    assert align_export.windows_path(Path("/mnt/c/Users/jij/OneDrive")) == r"C:\Users\jij\OneDrive"
    assert align_export.windows_path(Path("/home/jij")) == "/home/jij"
    if Path("/mnt").is_dir():
        assert align_export.to_path(r"C:\Users\jij") == Path("/mnt/c/Users/jij")


def data_folders(html: str) -> list[str]:
    """Every `data-folder` value, as a browser reads it."""
    from html.parser import HTMLParser

    found: list[str] = []

    class Reader(HTMLParser):
        def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
            found.extend(v or "" for k, v in attrs if k == "data-folder")

    Reader().feed(html)
    return found


def test_the_dialog_routes(app: TestClient, tmp_path: Path) -> None:
    quoted = tmp_path / """met "aanhalingstekens" en 'enkele'"""
    quoted.mkdir()
    html = app.get("/mappen", params={"folder": str(tmp_path)}).text
    assert 'id="mapkiezer"' in html and "Kies deze map" in html
    # The chosen path rides in a data attribute — escaped there, read back whole — never in a script.
    assert data_folders(app.get("/mappen", params={"folder": str(quoted)}).text) == [str(quoted)]
    made = app.post("/map_maken", data={"folder": str(tmp_path), "naam": "backup"})
    assert made.status_code == 200 and (tmp_path / "backup").is_dir()
    assert app.get("/map_maken", params={"folder": str(tmp_path), "naam": "x"}).status_code == 405
    page = app.get("/opslaan").text
    assert "Bladeren…" in page and 'id="mapkiezer-dialog"' in page


# --- the autosave switch ---------------------------------------------------------------------


def checked(html: str) -> bool:
    """Whether the page's autosave switch is on, as the browser would draw it."""
    switch = re.search(r'<input[^>]*name="autosave"[^>]*>', html)
    assert switch, "no autosave switch on the page"
    return " checked" in switch.group(0)


def test_the_switch_stays_on_once_turned_on(app: TestClient, tmp_path: Path) -> None:
    out = tmp_path / "eigen"
    out.mkdir()
    app.post("/opslaan_instellen", data={"format": "json", "folder": str(out)})
    assert not checked(app.get("/opslaan").text)
    reply = app.post("/autosave_zetten", data={"autosave": "on", "minutes": "15"})
    assert checked(reply.text) and reply.headers["HX-Trigger"] == "gewijzigd"
    # Stored, so the settings page draws it on when it is opened again.
    assert checked(app.get("/opslaan").text)
    saving = align.State.load(tmp_path / "state.json").saving
    assert saving.autosave and saving.minutes == 15
    # Saving the folder again leaves it on: the form does not carry the switch.
    app.post("/opslaan_instellen", data={"format": "csv", "folder": str(out)})
    assert checked(app.get("/opslaan").text)
    app.post("/autosave_zetten", data={"minutes": "15"})  # an unticked checkbox sends nothing
    assert not checked(app.get("/opslaan").text)


def test_the_switch_waits_for_a_folder(app: TestClient, tmp_path: Path) -> None:
    reply = app.post("/autosave_zetten", data={"autosave": "on"})
    assert not checked(reply.text) and "kies eerst een map" in reply.text
    assert not align.State.load(tmp_path / "state.json").saving.autosave


def test_turning_it_on_saves_what_is_unsaved(app: TestClient, tmp_path: Path) -> None:
    out = tmp_path / "eigen"
    out.mkdir()
    app.post("/opslaan_instellen", data={"format": "json", "folder": str(out)})
    app.post("/bewaar", data={"key": "DPR-227", "note": "na de laatste kopie"})
    assert align.State.load(tmp_path / "state.json").unsaved
    app.post("/autosave_zetten", data={"autosave": "on"})
    assert not align.State.load(tmp_path / "state.json").unsaved
    assert "na de laatste kopie" in (out / "afstemming-kwartaalplanning.json").read_text()


def test_saving_is_a_button_in_the_blue_bar_not_a_tab(app: TestClient) -> None:
    html = app.get("/").text
    tabs = re.search(r'<nav[^>]*class="tabs".*?</nav>', html, re.S)
    assert tabs and "Opslaan" not in tabs.group(0)
    bar = re.search(r'<header class="brandbar">.*?</header>', html, re.S)
    assert bar and 'href="/opslaan"' in bar.group(0) and "Opslaan/Openen" in bar.group(0)
    assert 'id="autosave"' not in html, "the autosave switch lives on the Opslaan/Openen page"
    settings = app.get("/opslaan").text
    assert 'id="autosave"' in settings
    assert re.search(r'<a [^>]*aria-current="page"[^>]*>Opslaan/Openen</a>', settings)
    assert "Terug naar de vergelijking" in settings


# --- loading a copy back -----------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["xlsx", "ods", "json"])
def test_a_copy_loads_back_whole(tmp_path: Path, fmt: str) -> None:
    state, rows = saved_state(tmp_path)
    loaded = align_export.load_copy(fmt, align_export.render(fmt, rows, state), rows)
    assert loaded.complete and not loaded.skipped
    assert loaded.state.decisions == {"DPR-227": align.Decision(wanted=6, note="Dubbel; zie actie")}
    assert loaded.state.links == {"e5": ""}
    [action] = loaded.state.actions
    assert (action.key, action.text, action.who, action.status) == ("DPR-227", "Navragen", "Agnes Dubbink", "open")
    assert action.created == state.actions[0].created


def test_a_csv_copy_brings_back_the_decisions_only(tmp_path: Path) -> None:
    state, rows = saved_state(tmp_path)
    state.decisions["DPR-227"] = align.Decision(wanted=7.5, note="komma")
    loaded = align_export.load_copy("csv", align_export.render("csv", rows, state), rows)
    assert not loaded.complete and loaded.state.decisions["DPR-227"] == align.Decision(wanted=7.5, note="komma")
    assert loaded.counts() == "1 besluiten"


def test_a_copy_without_keys_is_matched_on_dpr_or_name(tmp_path: Path) -> None:
    _, rows = saved_state(tmp_path)
    other = next(r for r in rows if r.key != "DPR-227" and r.monday)
    table = [
        ["DPR", "Epic (monday.com)", "Gewenst STP", "Toelichting"],
        ["DPR-227", "", 4, ""],
        ["", other.name, None, "op naam"],
        ["DPR-999", "Bestaat niet", 1, ""],
    ]
    loaded = align_export._from_tables({"Afstemming": table}, rows)
    assert loaded.state.decisions["DPR-227"].wanted == 4
    assert loaded.state.decisions[other.key].note == "op naam"
    assert loaded.skipped == ["DPR-999"]


def test_loading_a_copy_in_the_app(app: TestClient, tmp_path: Path) -> None:
    state, rows = saved_state(tmp_path)
    data = align_export.render("xlsx", rows, state)
    app.post("/bewaar", data={"key": "DPR-227", "wanted": "1", "note": "wordt vervangen"})
    reply = app.post("/laden", files={"bestand": ("kopie.xlsx", data)})
    assert "Geladen uit kopie.xlsx" in reply.text and "1 besluiten, 1 acties, 1 koppelingen" in reply.text
    assert "gewijzigd" in reply.headers["HX-Trigger"]
    now = align.State.load(tmp_path / "state.json")
    assert now.decision("DPR-227").note == "Dubbel; zie actie" and now.links == {"e5": ""}
    [backup] = tmp_path.glob("state_voor-laden_*.json")
    assert json.loads(backup.read_text())["decisions"]["DPR-227"]["note"] == "wordt vervangen"


def test_a_bad_file_is_a_message(app: TestClient, tmp_path: Path) -> None:
    assert "Laden lukte niet" in app.post("/laden", files={"bestand": ("kopie.xlsx", b"geen zip")}).text
    assert "geen kopie" in app.post("/laden", files={"bestand": ("foto.png", b"x")}).text
    assert not (tmp_path / "state.json").exists()


def test_loading_the_copy_in_the_folder(app: TestClient, tmp_path: Path) -> None:
    folder = tmp_path / "kopie"
    folder.mkdir()
    state, rows = saved_state(tmp_path)
    (folder / f"{align_export.FILE_STEM}.json").write_bytes(align_export.render("json", rows, state))
    with align.editing(tmp_path / "state.json") as now:
        now.saving.folder = str(folder)
        now.decisions = {}
    assert f"Laad {align_export.FILE_STEM}.json" in app.get("/opslaan").text
    assert "Geladen uit" in app.post("/laden_uit_map", data={"fmt": "json"}).text
    assert align.State.load(tmp_path / "state.json").decision("DPR-227").wanted == 6
    assert "niet (meer)" in app.post("/laden_uit_map", data={"fmt": "ods"}).text
