"""The alignment's decisions, as a file of the user's own — temporary, like `align.py`.

The decisions file lives in `docs/tmp`, which git ignores: nothing reaches the public
repository, and nothing is backed up either. This module writes a copy the user keeps
where they choose — a OneDrive folder, say — as Excel, LibreOffice Calc, CSV or JSON.

- **Two tables.** *Afstemming* is one row per comparison row: both sides' STP, the
  splits, the difference and the decision. *Acties* is one row per action. Excel and Calc
  get them as two sheets; CSV, being one table, carries the actions as a column of the
  first; JSON carries the decisions file itself plus the first table.
- **CSV is written for a Dutch Excel**: `;` between fields and a BOM, so it opens as
  columns rather than as one long line.
- **Never into the repository.** A folder inside the checkout is refused, because that is
  exactly where the copy must not end up. A Windows path (`C:\\Users\\…`) is read as its
  WSL mount (`/mnt/c/Users/…`), since that is how this runs on Jelle's machine.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
from dataclasses import dataclass, replace
from datetime import date, datetime
from functools import cache
from pathlib import Path
from typing import Any

from mondaycom import align, planning
from mondaycom.align import Row, State
from mondaycom.config import DISCIPLINES

#: The copy's file name, without its extension. One file, rewritten on every save.
FILE_STEM = "afstemming-kwartaalplanning"

MEDIA_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "ods": "application/vnd.oasis.opendocument.spreadsheet",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
}

Cell = str | float | int | None


# --- the folder -------------------------------------------------------------------------


def to_path(text: str) -> Path:
    """The folder as typed, as a path here: quotes stripped, `~` expanded, and a Windows path
    (`C:\\Users\\…`) read as its WSL mount (`/mnt/c/Users/…`). Not checked for existence."""
    text = text.strip().strip('"').strip()
    drive = re.match(r"^([A-Za-z]):[\\/]*(.*)$", text)
    if drive and Path("/mnt").is_dir():
        return Path("/mnt") / drive[1].lower() / drive[2].replace("\\", "/")
    return Path(text).expanduser()


def windows_path(path: Path) -> str:
    """`/mnt/c/Users/jij` as a Windows user knows it, `C:\\Users\\jij`; any other path as it is."""
    parts = path.parts
    if len(parts) >= 3 and parts[:2] == ("/", "mnt") and len(parts[2]) == 1:
        return f"{parts[2].upper()}:\\" + "\\".join(parts[3:])
    return str(path)


def unusable(path: Path) -> str:
    """Why the copy cannot go into this existing folder, in Dutch; empty when it can."""
    real = path.resolve()
    if real == align.REPO or align.REPO in real.parents:
        return f"{path} zit in de repository, en die is openbaar. Kies een map erbuiten, bijvoorbeeld in OneDrive."
    if not os.access(real, os.W_OK):
        return f"In {path} mag niet geschreven worden."
    return ""


def resolve_folder(text: str) -> Path:
    """The folder the user typed, as a path here — or a `ValueError` saying what is wrong."""
    if not text.strip():
        raise ValueError("Kies eerst een map.")
    path = to_path(text)
    if not path.is_absolute():
        raise ValueError(f"Geef een volledig pad, niet {text.strip()!r}.")
    if not path.is_dir():
        raise ValueError(f"De map {path} bestaat niet. Kies of maak hem met “Bladeren…”.")
    problem = unusable(path)
    if problem:
        raise ValueError(problem)
    return path


# --- browsing for a folder ----------------------------------------------------------------
# The Opslaan tab's "Bladeren…" dialog. The app writes the copy, so it is this machine's
# folders that matter, not the browser's: the dialog lists them from here.

#: Folders never offered: hidden ones, and Windows' own bookkeeping.
SKIPPED_PREFIXES = (".", "$")
SKIPPED_NAMES = frozenset({"System Volume Information", "AppData", "Application Data", "Local Settings"})

#: Enough for any folder a person keeps documents in; a system folder is cut off here.
MAX_FOLDERS = 400


@dataclass(frozen=True)
class Listing:
    """One folder as the dialog shows it."""

    path: Path
    folders: list[Path]
    #: Why the copy cannot go here; empty when it can.
    problem: str = ""
    #: The first folder of a typed path that does not exist yet — offered as "Nieuwe map".
    missing: str = ""
    #: Why the folder could not be read, or a new one not made.
    error: str = ""
    #: More subfolders than `MAX_FOLDERS`.
    cut: bool = False

    @property
    def parent(self) -> Path | None:
        return self.path.parent if self.path.parent != self.path else None


def start_folder(text: str) -> tuple[Path, str]:
    """Where the dialog opens for what is typed: that folder, or — when it does not exist —
    the nearest one above it that does, and the name of the first missing folder below it.
    Nothing typed (or nothing usable) opens the first suggestion, else home."""
    path = to_path(text) if text.strip() else None
    if path is None or not path.is_absolute():
        first = suggested_folders()
        return (Path(first[0]) if first else Path.home()), ""
    missing = ""
    while not path.is_dir() and path.parent != path:
        missing, path = path.name, path.parent
    return path, missing


def browse(text: str) -> Listing:
    """The subfolders of the folder `text` means, as `start_folder` finds it."""
    path, missing = start_folder(text)
    try:
        entries = sorted(
            (e for e in os.scandir(path) if not e.name.startswith(SKIPPED_PREFIXES) and e.name not in SKIPPED_NAMES),
            key=lambda e: e.name.casefold(),
        )
        folders = [Path(e.path) for e in entries if _is_dir(e)]
    except OSError as exc:
        return Listing(path, [], unusable(path), missing, error=f"{path} kan niet worden gelezen: {exc.strerror}")
    return Listing(path, folders[:MAX_FOLDERS], unusable(path), missing, cut=len(folders) > MAX_FOLDERS)


def _is_dir(entry: os.DirEntry[str]) -> bool:
    try:
        return entry.is_dir()
    except OSError:  # a Windows junction the mount will not follow
        return False


#: What a folder name may not hold, on Windows or here.
BAD_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def make_folder(parent: str, name: str) -> Listing:
    """Make `name` inside `parent` and list it — or list `parent` with the reason it failed."""
    name = name.strip()
    base = to_path(parent)
    problem = ""
    if not name or name in (".", "..") or BAD_NAME.search(name) or len(name) > 120:
        problem = f'{name!r} kan geen mapnaam zijn: geen / \\ : * ? " < > |, en niet leeg.'
    elif not base.is_absolute() or not base.is_dir():
        problem = f"De map {base} bestaat niet."
    elif blocked := unusable(base):
        problem = blocked
    else:
        try:
            (base / name).mkdir()
        except FileExistsError:
            pass  # there already: just go in
        except OSError as exc:
            problem = f"{name} kon niet worden gemaakt: {exc.strerror}"
    if problem:
        return replace(browse(str(base)), error=problem)
    return browse(str(base / name))


@cache
def suggested_folders() -> list[str]:
    """Folders worth offering: home and, under WSL, the Windows user's OneDrive and Documents.

    Asked once per run: scanning `/mnt/c` goes over WSL's 9P mount, and folders do not come
    and go while the app runs.
    """
    found = [Path.home() / "Documents", Path.home()]
    users = Path("/mnt/c/Users")
    if users.is_dir():
        for user in sorted(users.iterdir()):
            if user.name in ("Public", "Default", "Default User", "All Users"):
                continue
            onedrive = [d for d in sorted(user.glob("OneDrive*")) if "Temp" not in d.name]  # not OneDriveCloudTemp
            found += [*onedrive, user / "Documents", user / "Desktop"]
    return [str(p) for p in found if _writable_dir(p)]


def _writable_dir(path: Path) -> bool:
    try:
        return path.is_dir() and os.access(path, os.W_OK)
    except OSError:  # a WSL mount can refuse even the question
        return False


# --- the tables -------------------------------------------------------------------------


def _number(value: float | None) -> Cell:
    if value is None or value == float("inf") or value == float("-inf"):
        return None
    return round(value, 2)


def comparison_table(rows: list[Row], state: State) -> tuple[list[str], list[list[Cell]]]:
    """One line per row, every number the page shows plus what was decided."""
    head = [
        "DPR",
        "Epic (monday.com)",
        "Epic (database)",
        "Soort",
        "Groep",
        "Laag",
        "Database STP",
        *[f"Database {align.SHORT[d]}" for d in DISCIPLINES],
        "Verdeling database",
        "monday.com STP-TODO",
        "Verdeling monday.com",
        "Verschil",
        "Verschil %",
        "Meer dan 10% af",
        "Gewenst STP",
        "Waar aanpassen",
        "Toelichting",
        "Open acties",
    ]
    lines: list[list[Cell]] = []
    for r in rows:
        d = state.decision(r.key)
        points = r.board_points
        ratio = r.ratio
        open_actions = sum(a.status != "klaar" for a in state.actions_for(r.key))
        lines.append(
            [
                r.dpr,
                r.monday.name if r.monday else "",
                " + ".join(e.name for e in r.board),
                align.KINDS_NL[r.kind],
                r.monday.group_title if r.monday else "",
                planning.LAYERS_NL[r.layer] if r.layer else "",
                _number(r.board_total) if r.board else None,
                *[_number(points.get(dis, 0.0)) if r.board else None for dis in DISCIPLINES],
                align.split_text(r.board_split) if r.board else "",
                _number(r.monday_total),
                align.split_text(r.monday_split) if r.monday else "",
                _number(r.diff),
                _number(ratio * 100) if ratio is not None else None,
                "ja" if r.flagged else "",
                _number(d.wanted),
                align.WHERE[d.where] if d.where else "",
                d.note,
                open_actions or None,
            ]
        )
    return head, lines


def actions_table(rows: list[Row], state: State) -> tuple[list[str], list[list[Cell]]]:
    head = ["DPR", "Epic", "Wie", "Actie", "Status", "Aangemaakt", "Verstuurd", "Klaar"]
    by_key = {r.key: r for r in rows}
    lines: list[list[Cell]] = []
    for a in state.actions:
        r = by_key.get(a.key)
        lines.append([r.dpr if r else "", r.name if r else a.key, a.who, a.text, a.status, a.created, a.sent, a.done])
    return head, lines


def _actions_text(state: State, key: str) -> str:
    return "; ".join(f"{a.who or 'niemand'}: {a.text} ({a.status})" for a in state.actions_for(key))


# --- the formats ------------------------------------------------------------------------


def render(fmt: str, rows: list[Row], state: State) -> bytes:
    """The copy, as the bytes of one file in `fmt`."""
    if fmt in ("xlsx", "ods"):
        sheets = {"Afstemming": comparison_table(rows, state), "Acties": actions_table(rows, state)}
        return _xlsx(sheets) if fmt == "xlsx" else _ods(sheets)
    if fmt == "csv":
        return _csv(rows, state, comparison_table(rows, state))
    if fmt == "json":
        head, lines = comparison_table(rows, state)
        data = {
            "opgeslagen": datetime.now().isoformat(timespec="seconds"),
            **state.content(),
            "afstemming": [dict(zip(head, line, strict=True)) for line in lines],
        }
        return json.dumps(data, indent=2, ensure_ascii=False).encode("utf-8")
    raise ValueError(f"Onbekend formaat {fmt!r}; kies uit {', '.join(align.FORMATS)}.")


def _csv(rows: list[Row], state: State, table: tuple[list[str], list[list[Cell]]]) -> bytes:
    head, lines = table
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow([*head, "Acties"])
    for r, line in zip(rows, lines, strict=True):
        cells = [str(c).replace(".", ",") if isinstance(c, float) else c for c in line]
        writer.writerow([*cells, _actions_text(state, r.key)])
    return ("\ufeff" + out.getvalue()).encode("utf-8")


def _xlsx(sheets: dict[str, tuple[list[str], list[list[Cell]]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)  # type: ignore[arg-type]
    for title, (head, lines) in sheets.items():
        sheet = book.create_sheet(title)
        sheet.append(head)
        for line in lines:
            sheet.append(line)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        if lines:
            sheet.auto_filter.ref = sheet.dimensions
        for i, name in enumerate(head, start=1):
            longest = max([len(name), *(len(str(line[i - 1] or "")) for line in lines)])
            sheet.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 60)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _ods(sheets: dict[str, tuple[list[str], list[list[Cell]]]]) -> bytes:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.style import Style, TextProperties
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    doc = OpenDocumentSpreadsheet()
    bold = Style(name="kop", family="table-cell")
    bold.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(bold)

    def cell(value: Cell, style: Any = None) -> Any:
        kw = {"stylename": style} if style is not None else {}
        if isinstance(value, int | float):
            c = TableCell(valuetype="float", value=value, **kw)
        else:
            c = TableCell(valuetype="string", **kw)
        if value is not None and value != "":
            c.addElement(P(text=str(value)))
        return c

    for title, (head, lines) in sheets.items():
        table = Table(name=title)
        top = TableRow()
        for name in head:
            top.addElement(cell(name, bold))
        table.addElement(top)
        for line in lines:
            tr = TableRow()
            for value in line:
                tr.addElement(cell(value))
            table.addElement(tr)
        doc.spreadsheet.addElement(table)
    out = io.BytesIO()
    doc.write(out)
    return out.getvalue()


# --- saving -----------------------------------------------------------------------------


def save_copy(state: State, rows: list[Row]) -> Path:
    """Write the copy to the folder and format in `state.saving`, and record that it did.

    Written to a temporary name and moved over the old copy, so a crash halfway never
    leaves a broken file where the good one was. The caller saves `state` afterwards, and
    records an error that this raises.
    """
    saving = state.saving
    folder = resolve_folder(saving.folder)
    target = folder / f"{FILE_STEM}.{saving.format}"
    tmp = folder / f".{target.name}.tmp"
    tmp.write_bytes(render(saving.format, rows, state))
    tmp.replace(target)
    record_saved(state, str(target))
    return target


def record_saved(state: State, where: str) -> None:
    """Mark the current decisions as saved — to a folder, or as a download."""
    state.saving.saved_at = datetime.now().isoformat(timespec="seconds")
    state.saving.saved_to = where
    state.saving.saved_print = state.fingerprint
    state.saving.error = ""


def autosave_due(state: State, now: datetime | None = None) -> bool:
    """Autosave is on, a folder is set, something changed, and the interval has passed."""
    saving = state.saving
    if not (saving.autosave and saving.folder and state.unsaved):
        return False
    if not saving.saved_at:
        return True
    last = datetime.fromisoformat(saving.saved_at)
    return ((now or datetime.now()) - last).total_seconds() >= saving.minutes * 60


def download_name(fmt: str) -> str:
    return f"{FILE_STEM}-{date.today().isoformat()}.{fmt}"
